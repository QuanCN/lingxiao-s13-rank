# -*- coding: utf-8 -*-
"""从 data.xlsx 生成 index.html（属性排名，不含爵位/阵容/集结进分）"""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("缺少 openpyxl，请先: pip install openpyxl")

ROOT = Path(__file__).resolve().parent
DATA_XLSX = ROOT / "data.xlsx"
OUT_HTML = ROOT / "index.html"

W_INF, W_ARC = 0.55, 0.45


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('"', "").replace("\n", "").replace("\r", "")
    if not s or s.lower() == "none":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def to_str(v):
    if v is None:
        return ""
    return str(v).strip().replace("\n", "").replace("\r", "").replace('"', "")


def to_bool_gong(v):
    s = to_str(v)
    if s in ("是", "Y", "y", "true", "True", "1"):
        return True
    if s in ("否", "N", "n", "false", "False", "0"):
        return False
    return None


def esc(s):
    return (
        str(s)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def fmt(n, d=2):
    if n is None:
        return "-"
    if isinstance(n, float) and n == int(n):
        return str(int(n))
    if isinstance(n, (int, float)):
        s = f"{n:.{d}f}".rstrip("0").rstrip(".")
        return s
    return str(n)


def load_members(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    members = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        name = to_str(row[1] if len(row) > 1 else None)
        if not name:
            continue
        members.append(
            {
                "seq": to_num(row[0]),
                "name": name,
                "title": to_num(row[2]) if len(row) > 2 else None,
                "faction": to_str(row[3]) if len(row) > 3 else "",
                "gong4": to_bool_gong(row[4]) if len(row) > 4 else None,
                "rally": to_num(row[5]) if len(row) > 5 else None,
                "inf_hp": to_num(row[6]) if len(row) > 6 else None,
                "inf_def": to_num(row[7]) if len(row) > 7 else None,
                "arc_atk": to_num(row[8]) if len(row) > 8 else None,
                "arc_dmg": to_num(row[9]) if len(row) > 9 else None,
            }
        )
    return members, ws.title


def score_members(members):
    complete, incomplete = [], []
    for m in members:
        attrs = [m["inf_hp"], m["inf_def"], m["arc_atk"], m["arc_dmg"]]
        if all(a is not None for a in attrs):
            m["inf_sum"] = m["inf_hp"] + m["inf_def"]
            m["arc_sum"] = m["arc_atk"] + m["arc_dmg"]
            complete.append(m)
        else:
            issues = []
            for k, lab in (
                ("inf_hp", "缺步兵生命"),
                ("inf_def", "缺步兵防御"),
                ("arc_atk", "缺弓兵攻击"),
                ("arc_dmg", "缺弓兵破坏"),
            ):
                if m[k] is None:
                    issues.append(lab)
            m["issues"] = issues
            incomplete.append(m)

    if not complete:
        raise SystemExit("没有四项兵种齐全的成员，无法生成排名")

    max_inf = max(m["inf_sum"] for m in complete)
    max_arc = max(m["arc_sum"] for m in complete)
    for m in complete:
        inf_n = m["inf_sum"] / max_inf * 100
        arc_n = m["arc_sum"] / max_arc * 100
        m["score"] = round(W_INF * inf_n + W_ARC * arc_n, 2)

    complete.sort(key=lambda x: (-x["score"], -x["inf_sum"], -x["arc_sum"], x["name"]))
    for i, m in enumerate(complete, 1):
        m["rank"] = i

    blank, partial = [], []
    for m in incomplete:
        has_combat = any(m[k] is not None for k in ("inf_hp", "inf_def", "arc_atk", "arc_dmg"))
        if not has_combat and m["rally"] is None and m["title"] is None:
            blank.append(m)
        else:
            partial.append(m)

    meta = {
        "source": DATA_XLSX.name,
        "total": len(members),
        "complete": len(complete),
        "incomplete": len(incomplete),
        "avgScore": round(sum(m["score"] for m in complete) / len(complete), 2),
        "avgInf": round(sum(m["inf_sum"] for m in complete) / len(complete), 1),
        "avgArc": round(sum(m["arc_sum"] for m in complete) / len(complete), 1),
        "faction": dict(Counter(m["faction"] for m in members if m["faction"])),
        "gong": {
            "yes": sum(1 for m in members if m["gong4"] is True),
            "no": sum(1 for m in members if m["gong4"] is False),
            "unk": sum(1 for m in members if m["gong4"] is None),
        },
        "max_inf": max_inf,
        "max_arc": max_arc,
        "formula": f"属性分 = {W_INF}*步兵归一 + {W_ARC}*弓兵归一（不含爵位/阵容/集结）",
        "inf_rule": f"步兵归一 = (生命+防御)/{max_inf}*100",
        "arc_rule": f"弓兵归一 = (攻击+破坏)/{max_arc}*100",
    }
    return complete, blank, partial, meta


def render_html(complete, blank, partial, meta, members):
    faction = meta["faction"]
    faction_total = sum(faction.values()) or 1
    order = [("魏", "#4a7fd4"), ("蜀", "#3d9a6a"), ("吴", "#c47a2c"), ("群", "#8b6bc7")]
    acc = 0.0
    stops = []
    for name, color in order:
        v = faction.get(name, 0)
        if not v:
            continue
        start_p = acc
        acc += v / faction_total * 100
        stops.append(f"{color} {start_p:.2f}% {acc:.2f}%")
    pie_css = ", ".join(stops) if stops else "#ccc 0 100%"

    max_score = complete[0]["score"]
    bars = []
    for m in complete[:10]:
        pct = round(m["score"] / max_score * 100, 1)
        bars.append(
            f'<div class="bar-row"><div class="bar-label">{esc(m["name"])}</div>'
            f'<div class="bar-track"><div class="bar-fill" style="width:{pct}%"></div></div>'
            f'<div class="bar-val">{fmt(m["score"])}</div></div>'
        )

    rank_rows = []
    for m in complete:
        gong = "是" if m["gong4"] is True else ("否" if m["gong4"] is False else "-")
        tone = " top" if m["rank"] <= 10 else (" warn" if m["gong4"] is False else "")
        tag = ' <span class="tag">宫四否</span>' if m["gong4"] is False else ""
        rank_rows.append(
            f'<tr class="{tone.strip()}"><td class="num">{m["rank"]}</td>'
            f'<td>{esc(m["name"])}{tag}</td><td>{esc(m["faction"] or "-")}</td>'
            f'<td class="num">{fmt(m["title"], 0)}</td><td class="num">{fmt(m["rally"])}</td>'
            f'<td class="num">{fmt(m["inf_sum"], 1)}</td><td class="num">{fmt(m["arc_sum"], 1)}</td>'
            f'<td class="num score">{fmt(m["score"])}</td><td>{gong}</td></tr>'
        )

    top_inf = sorted(complete, key=lambda x: -x["inf_sum"])[:10]
    top_arc = sorted(complete, key=lambda x: -x["arc_sum"])[:10]
    inf_rows = "".join(
        f"<tr><td>{i}</td><td>{esc(m['name'])}</td><td>{esc(m['faction'] or '-')}</td>"
        f"<td class=num>{fmt(m['inf_sum'],1)}</td><td class=num>{fmt(m['inf_hp'],1)}</td>"
        f"<td class=num>{fmt(m['inf_def'],1)}</td></tr>"
        for i, m in enumerate(top_inf, 1)
    )
    arc_rows = "".join(
        f"<tr><td>{i}</td><td>{esc(m['name'])}</td><td>{esc(m['faction'] or '-')}</td>"
        f"<td class=num>{fmt(m['arc_sum'],1)}</td><td class=num>{fmt(m['arc_atk'],1)}</td>"
        f"<td class=num>{fmt(m['arc_dmg'],1)}</td></tr>"
        for i, m in enumerate(top_arc, 1)
    )
    partial_rows = "".join(
        f"<tr><td class=num>{fmt(m['seq'],0)}</td><td>{esc(m['name'])}</td>"
        f"<td>{esc('；'.join(m.get('issues') or []))}</td>"
        f"<td class=num>{fmt(m['inf_hp'],1)}</td><td class=num>{fmt(m['inf_def'],1)}</td>"
        f"<td class=num>{fmt(m['arc_atk'],1)}</td><td class=num>{fmt(m['arc_dmg'],1)}</td>"
        f"<td>{'是' if m['gong4'] is True else ('否' if m['gong4'] is False else '-')}</td></tr>"
        for m in partial
    )
    gong_no = [m for m in members if m["gong4"] is False]
    gong_no_text = "、".join(esc(m["name"]) + f'（集结 {fmt(m["rally"])}）' for m in gong_no) or "-"
    pie_legend = "".join(
        f'<div class="legend"><span class="dot f-{esc(k)}"></span>{esc(k)} {v}'
        f"（{v * 100 / faction_total:.0f}%）</div>"
        for k, v in faction.items()
    )
    faction_opts = "".join(f'<option value="{esc(k)}">{esc(k)}</option>' for k in faction)

    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>凌霄 S13 属性排名</title>
<style>
  :root {{
    --bg: #0f1115; --panel: #171a21; --border: #2a2f3a; --text: #e8eaed;
    --muted: #9aa3b2; --dim: #6b7380; --accent: #5b8def; --warn: #c9a227;
    --row2: #151820;
  }}
  * {{ box-sizing: border-box; }}
  body {{ margin:0; font-family:"Segoe UI","PingFang SC","Microsoft YaHei",sans-serif; background:var(--bg); color:var(--text); line-height:1.5; }}
  .wrap {{ max-width:1100px; margin:0 auto; padding:28px 20px 60px; }}
  h1 {{ font-size:26px; font-weight:650; margin:0 0 6px; }}
  h2 {{ font-size:18px; font-weight:600; margin:28px 0 10px; }}
  h3 {{ font-size:15px; font-weight:600; margin:18px 0 8px; color:var(--muted); }}
  .sub {{ color:var(--muted); font-size:13px; margin-bottom:22px; }}
  .stats {{ display:grid; grid-template-columns:repeat(4,1fr); gap:12px; margin-bottom:18px; }}
  .stat {{ background:var(--panel); border:1px solid var(--border); border-radius:8px; padding:14px 16px; }}
  .stat .v {{ font-size:22px; font-weight:650; }}
  .stat .l {{ font-size:12px; color:var(--muted); margin-top:4px; }}
  .callout {{ background:#152033; border:1px solid #2a4060; border-radius:8px; padding:12px 14px; font-size:13px; color:#c5d4ea; margin-bottom:22px; }}
  .callout strong {{ color:var(--text); }}
  .grid2 {{ display:grid; grid-template-columns:1.3fr 1fr; gap:16px; }}
  .panel {{ background:var(--panel); border:1px solid var(--border); border-radius:8px; padding:14px 16px; }}
  .bar-row {{ display:grid; grid-template-columns:120px 1fr 48px; gap:8px; align-items:center; margin:6px 0; font-size:12px; }}
  .bar-label {{ overflow:hidden; text-overflow:ellipsis; white-space:nowrap; color:var(--muted); }}
  .bar-track {{ height:10px; background:#232833; border-radius:4px; overflow:hidden; }}
  .bar-fill {{ height:100%; background:var(--accent); border-radius:4px; }}
  .bar-val {{ text-align:right; font-variant-numeric:tabular-nums; }}
  .pie-wrap {{ display:flex; gap:20px; align-items:center; padding:8px 0; }}
  .pie {{ width:140px; height:140px; border-radius:50%; background:conic-gradient({pie_css}); flex-shrink:0; }}
  .legend {{ font-size:13px; color:var(--muted); margin:6px 0; }}
  .dot {{ display:inline-block; width:10px; height:10px; border-radius:50%; margin-right:8px; }}
  .f-魏 {{ background:#4a7fd4; }} .f-蜀 {{ background:#3d9a6a; }} .f-吴 {{ background:#c47a2c; }} .f-群 {{ background:#8b6bc7; }}
  .tools {{ display:flex; flex-wrap:wrap; gap:8px; margin:10px 0 12px; align-items:center; }}
  .tools label {{ font-size:12px; color:var(--muted); }}
  select, input[type="search"] {{ background:#1c2028; color:var(--text); border:1px solid var(--border); border-radius:6px; padding:6px 10px; font-size:13px; }}
  table {{ width:100%; border-collapse:collapse; font-size:13px; background:var(--panel); border:1px solid var(--border); border-radius:8px; overflow:hidden; }}
  th, td {{ padding:8px 10px; border-bottom:1px solid var(--border); text-align:left; }}
  th {{ background:#1a1e27; color:var(--muted); font-weight:600; font-size:12px; position:sticky; top:0; z-index:1; }}
  tr:nth-child(even) td {{ background:var(--row2); }}
  tr.top td {{ color:#d8e6ff; }} tr.warn td {{ color:#e6d28a; }}
  td.num, th.num {{ text-align:right; font-variant-numeric:tabular-nums; }}
  td.score {{ font-weight:650; color:var(--accent); }}
  .tag {{ display:inline-block; font-size:10px; color:var(--warn); border:1px solid #5a4a1a; border-radius:4px; padding:0 4px; margin-left:4px; }}
  .split {{ display:grid; grid-template-columns:1fr 1fr; gap:16px; }}
  .table-scroll {{ max-height:560px; overflow:auto; border-radius:8px; }}
  .muted {{ color:var(--muted); font-size:13px; }}
  .dim {{ color:var(--dim); font-size:12px; margin-top:8px; }}
  @media (max-width:800px) {{ .stats,.grid2,.split {{ grid-template-columns:1fr; }} .bar-row {{ grid-template-columns:90px 1fr 40px; }} }}
  @media print {{
    body {{ background:#fff; color:#111; }}
    .panel,.stat,.callout,table {{ background:#fff!important; border-color:#ccc!important; color:#111!important; }}
    th {{ background:#eee!important; color:#333!important; }}
    .table-scroll {{ max-height:none; overflow:visible; }}
  }}
</style>
</head>
<body>
<div class="wrap">
  <h1>凌霄 S13 属性排名</h1>
  <div class="sub">数据源：{esc(meta["source"])} · 四项齐全 {meta["complete"]} 人 / 全表 {meta["total"]} 人 · 评分不含爵位、阵容、集结 · 替换 data.xlsx 后重新构建即可更新</div>

  <div class="stats">
    <div class="stat"><div class="v">{meta["complete"]}</div><div class="l">四项齐全可排名</div></div>
    <div class="stat"><div class="v">{meta["incomplete"]}</div><div class="l">属性缺失未入榜</div></div>
    <div class="stat"><div class="v">{fmt(meta["avgScore"])}</div><div class="l">完整组平均属性分</div></div>
    <div class="stat"><div class="v">{fmt(meta["avgInf"],0)} / {fmt(meta["avgArc"],0)}</div><div class="l">平均步兵 / 弓兵合计</div></div>
  </div>

  <div class="callout">
    <strong>综合属性分公式</strong><br/>
    {esc(meta["formula"])}。{esc(meta["inf_rule"])}；{esc(meta["arc_rule"])}。
  </div>

  <div class="grid2">
    <div class="panel">
      <h2 style="margin-top:0">属性分 Top10</h2>
      {"".join(bars)}
    </div>
    <div class="panel">
      <h2 style="margin-top:0">阵营分布（仅参考）</h2>
      <div class="pie-wrap">
        <div class="pie"></div>
        <div>{pie_legend}
          <div class="dim">宫四兵：是 {meta["gong"]["yes"]} / 否 {meta["gong"]["no"]} / 未填 {meta["gong"]["unk"]}</div>
        </div>
      </div>
    </div>
  </div>

  <h2>完整属性排名</h2>
  <div class="tools">
    <label>阵营 <select id="faction"><option value="全部">全部</option>{faction_opts}</select></label>
    <label>排序 <select id="sort">
      <option value="score">属性分</option>
      <option value="inf">步兵合计</option>
      <option value="arc">弓兵合计</option>
      <option value="rally">集结参考</option>
      <option value="title">爵位参考</option>
    </select></label>
    <input type="search" id="q" placeholder="搜索成员名…" style="min-width:180px" />
  </div>
  <div class="table-scroll">
    <table id="rankTable">
      <thead><tr>
        <th class="num">#</th><th>成员</th><th>阵容</th><th class="num">爵位</th>
        <th class="num">集结</th><th class="num">步兵合计</th><th class="num">弓兵合计</th>
        <th class="num">属性分</th><th>宫四</th>
      </tr></thead>
      <tbody>{"".join(rank_rows)}</tbody>
    </table>
  </div>

  <div class="split" style="margin-top:24px">
    <div>
      <h2>步兵榜 Top10</h2>
      <table><thead><tr><th>#</th><th>成员</th><th>阵营</th><th class="num">合计</th><th class="num">生命</th><th class="num">防御</th></tr></thead>
      <tbody>{inf_rows}</tbody></table>
    </div>
    <div>
      <h2>弓兵榜 Top10</h2>
      <table><thead><tr><th>#</th><th>成员</th><th>阵营</th><th class="num">合计</th><th class="num">攻击</th><th class="num">破坏</th></tr></thead>
      <tbody>{arc_rows}</tbody></table>
    </div>
  </div>

  <h2>数据问题与未入榜</h2>
  <div class="callout" style="border-color:#5a4a1a;background:#2a2414;color:#e6d28a">
    <strong>宫四兵=否</strong><br/>{gong_no_text}
  </div>
  <h3>四项兵种缺项</h3>
  <table>
    <thead><tr><th class="num">序号</th><th>成员</th><th>问题</th><th class="num">步兵生命</th>
    <th class="num">步兵防御</th><th class="num">弓攻</th><th class="num">弓破</th><th>宫四</th></tr></thead>
    <tbody>{partial_rows}</tbody>
  </table>
  <h3>几乎空白</h3>
  <p class="muted">{"、".join(esc(m["name"]) for m in blank) or "-"}</p>
  <p class="dim">微信打开：本页链接 · 更新方式：替换仓库中的 data.xlsx 后自动重建</p>
</div>
<script>
(() => {{
  const factionSel = document.getElementById("faction");
  const sortSel = document.getElementById("sort");
  const q = document.getElementById("q");
  const tbody = document.querySelector("#rankTable tbody");
  const rows = Array.from(tbody.querySelectorAll("tr")).map((tr, i) => {{
    const tds = tr.children;
    return {{
      el: tr,
      name: tds[1].textContent,
      faction: tds[2].textContent,
      title: parseFloat(tds[3].textContent) || -1,
      rally: parseFloat(tds[4].textContent) || -1,
      inf: parseFloat(tds[5].textContent) || -1,
      arc: parseFloat(tds[6].textContent) || -1,
      score: parseFloat(tds[7].textContent) || -1,
      orig: i,
    }};
  }});
  function apply() {{
    const f = factionSel.value;
    const key = sortSel.value;
    const term = q.value.trim().toLowerCase();
    let list = rows.filter(r => {{
      if (f !== "全部" && r.faction !== f) return false;
      if (term && !r.name.toLowerCase().includes(term)) return false;
      return true;
    }});
    list.sort((a, b) => {{
      const av = a[key], bv = b[key];
      if (bv !== av) return bv - av;
      return a.orig - b.orig;
    }});
    tbody.innerHTML = "";
    list.forEach((r, i) => {{
      r.el.children[0].textContent = String(i + 1);
      tbody.appendChild(r.el);
    }});
  }}
  factionSel.addEventListener("change", apply);
  sortSel.addEventListener("change", apply);
  q.addEventListener("input", apply);
}})();
</script>
</body>
</html>
"""


def main():
    if not DATA_XLSX.exists():
        raise SystemExit(f"找不到 {DATA_XLSX}，请把表格命名为 data.xlsx 放在本目录")
    members, _sheet = load_members(DATA_XLSX)
    complete, blank, partial, meta = score_members(members)
    html = render_html(complete, blank, partial, meta, members)
    OUT_HTML.write_text(html, encoding="utf-8")
    print(f"OK -> {OUT_HTML}")
    print(f"ranked={meta['complete']} incomplete={meta['incomplete']} top={complete[0]['name']} {complete[0]['score']}")


if __name__ == "__main__":
    main()